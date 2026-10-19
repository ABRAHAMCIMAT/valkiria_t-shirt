"""Lectura, recuperación RAG y exportación del template T-Shirt."""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List

from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


@dataclass
class Role:
    name: str
    area: str = ""
    required: bool = True
    collaborators: float = 1
    cost_hourly: float = 0
    source_cost: float = 0
    source_basis: str = "mensual"
    source_row: int = 0


class WorkbookKnowledgeBase:
    """Índice local, auditable y sin servicios externos para el Excel."""

    def __init__(self, storage_dir: Path):
        self.storage_dir = Path(storage_dir)
        self.storage_dir.mkdir(parents=True, exist_ok=True)
        self.workbook_path = self.storage_dir / "template.xlsx"
        self.roles_path = self.storage_dir / "roles.json"
        self.index_path = self.storage_dir / "knowledge.json"
        self.chunks: List[Dict[str, Any]] = []
        self.vectorizer: TfidfVectorizer | None = None
        self.matrix = None

    def ingest(self, source: Path) -> Dict[str, Any]:
        shutil.copy2(source, self.workbook_path)
        workbook = load_workbook(self.workbook_path, data_only=True)
        chunks: List[Dict[str, Any]] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                values = [self._text(cell.value) for cell in row]
                populated = [value for value in values if value]
                if not populated:
                    continue
                text = " | ".join(populated)
                chunks.append({
                    "sheet": sheet.title,
                    "row": row[0].row,
                    "text": text[:8000],
                })
        self.chunks = chunks
        self._build_index()
        self.index_path.write_text(
            json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        roles = self._extract_roles(workbook)
        self.roles_path.write_text(
            json.dumps([asdict(role) for role in roles], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return {
            "sheets": workbook.sheetnames,
            "chunks": len(chunks),
            "roles": len(roles),
            "questionnaire_items": len(self.get_questionnaire()),
            "epics": len(self.get_epics()),
        }

    def load(self) -> bool:
        if not self.index_path.exists():
            return False
        self.chunks = json.loads(self.index_path.read_text(encoding="utf-8"))
        self._build_index()
        return True

    def search(self, query: str, limit: int = 6) -> List[Dict[str, Any]]:
        if not self.chunks:
            self.load()
        if not self.chunks or self.vectorizer is None:
            return []
        query_vector = self.vectorizer.transform([query])
        scores = cosine_similarity(query_vector, self.matrix)[0]
        indices = scores.argsort()[::-1][: max(1, min(limit, 20))]
        return [
            {**self.chunks[index], "score": round(float(scores[index]), 4)}
            for index in indices
            if scores[index] > 0
        ]

    def get_roles(self) -> List[Dict[str, Any]]:
        if not self.roles_path.exists():
            return []
        return json.loads(self.roles_path.read_text(encoding="utf-8"))

    def save_roles(self, roles: Iterable[Dict[str, Any]]) -> List[Dict[str, Any]]:
        normalized = []
        seen = set()
        for item in roles:
            name = str(item.get("name", "")).strip()
            if not name or name.casefold() in seen:
                continue
            seen.add(name.casefold())
            normalized.append({
                "name": name,
                "area": str(item.get("area", "")).strip(),
                "required": bool(item.get("required", True)),
                "collaborators": max(0, float(item.get("collaborators", 1))),
                "cost_hourly": max(0, float(item.get("cost_hourly", 0))),
                "source_cost": max(0, float(item.get("source_cost", 0))),
                "source_basis": str(item.get("source_basis", "hora")),
                "source_row": max(0, int(item.get("source_row", 0) or 0)),
            })
        self.roles_path.write_text(
            json.dumps(normalized, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        return normalized

    def get_questionnaire(self) -> List[Dict[str, Any]]:
        if not self.workbook_path.exists():
            return []
        ws = load_workbook(self.workbook_path, data_only=True)["Cuestionario"]
        questions: List[Dict[str, Any]] = []
        current = None
        for row in range(3, ws.max_row + 1):
            number = ws.cell(row, 2).value
            question = ws.cell(row, 4).value
            if number not in (None, ""):
                current = {
                    "number": number,
                    "row": row,
                    "topic": self._text(ws.cell(row, 3).value),
                    "question": self._text(question),
                    "options": [],
                }
                questions.append(current)
            if current and ws.cell(row, 5).value not in (None, ""):
                current["options"].append({
                    "row": row,
                    "answer": self._text(ws.cell(row, 5).value),
                    "size": self._text(ws.cell(row, 6).value),
                    "annual_cost": self._number(ws.cell(row, 7).value),
                })
        return questions

    def get_epics(self) -> List[Dict[str, Any]]:
        if not self.workbook_path.exists():
            return []
        ws = load_workbook(self.workbook_path, data_only=True)["Epic Backlog"]
        epics = []
        for row in range(3, ws.max_row + 1):
            name = self._text(ws.cell(row, 2).value)
            description = self._text(ws.cell(row, 3).value)
            if not name and not description:
                if epics:
                    break
                continue
            epics.append({
                "row": row,
                "number": ws.cell(row, 1).value,
                "name": name or f"Épica {ws.cell(row, 1).value}",
                "description": description,
                "acceptance_criteria": self._text(ws.cell(row, 4).value),
                "mvp": self._text(ws.cell(row, 5).value) or "MVP1",
            })
        return epics

    def export(self, output: Path, analysis: Dict[str, Any]) -> Path:
        """Llena una copia conservando formato, fórmulas e imágenes."""
        if not self.workbook_path.exists():
            raise FileNotFoundError("No se ha cargado un template")
        shutil.copy2(self.workbook_path, output)
        wb = load_workbook(output)

        role_rows = analysis.get("roles", [])
        ws_roles = wb["Roles"]
        ws_roles.cell(17, 10, "Costo por hora")
        existing_rows = {
            self._text(ws_roles.cell(row, 2).value).casefold(): row
            for row in range(1, ws_roles.max_row + 1)
            if self._text(ws_roles.cell(row, 2).value)
        }
        append_row = ws_roles.max_row + 2
        for role in role_rows:
            name = self._text(role.get("name"))
            row = int(role.get("source_row", 0) or 0) or existing_rows.get(name.casefold(), 0)
            if not row:
                row = append_row
                append_row += 1
            ws_roles.cell(row, 2, name)
            ws_roles.cell(row, 3, role.get("area"))
            ws_roles.cell(row, 4, "SI" if role.get("required", True) else "NO")
            ws_roles.cell(row, 5, role.get("collaborators", 1))
            ws_roles.cell(row, 10, role.get("cost_hourly", 0))

        questionnaire = analysis.get("questionnaire", [])
        ws_questionnaire = wb["Cuestionario"]
        for answer in questionnaire:
            row = int(answer.get("row", 0) or 0)
            if row > 0:
                ws_questionnaire.cell(row, 8, answer.get("answer", ""))
                ws_questionnaire.cell(row, 9, answer.get("cost", 0))

        ws_epics = wb["Epic Backlog"]
        for epic in analysis.get("epics", []):
            row = int(epic.get("row", 0) or 0)
            if row > 0:
                ws_epics.cell(row, 20, epic.get("size", "M"))
                ws_epics.cell(row, 21, epic.get("weeks", 0))
                ws_epics.cell(row, 29, epic.get("cost", 0))

        summary = analysis.get("summary", {})
        ws_result = wb["Resultado"]
        ws_result["B6"] = f"{summary.get('total_weeks', 0):.1f} Semanas"
        ws_result["B7"] = summary.get("total_cost", 0)
        ws_result["B8"] = summary.get("complexity", "M")
        ws_result["B9"] = summary.get("mvp_count", 1)
        ws_result["H10"] = summary.get("executive_summary", "")
        for row, item in zip(range(14, 20), analysis.get("additional_costs", [])):
            ws_result.cell(row, 1, item.get("concept", ""))
            ws_result.cell(row, 3, item.get("amount_mxn", 0))
            ws_result.cell(row, 4, item.get("amount_usd", 0))
            ws_result.cell(row, 5, item.get("description", ""))

        if "Arquitectura" in wb.sheetnames:
            wb["Arquitectura"]["B4"] = analysis.get("architecture", "Pendiente de validación arquitectónica")
        if "Proc de Neg (To-Be)" in wb.sheetnames:
            wb["Proc de Neg (To-Be)"]["B10"] = analysis.get("process", "Pendiente de validación del Product Owner")

        wb.save(output)
        return output

    def _extract_roles(self, workbook) -> List[Role]:
        ws = workbook["Roles"]
        roles: List[Role] = []
        for row in range(5, ws.max_row + 1):
            name = self._text(ws.cell(row, 2).value)
            if not name or name.casefold() in {
                "rol", "roles con costo", "roles sin costo", "fabrica de pruebas (qa)"
            }:
                continue
            area = self._text(ws.cell(row, 3).value)
            if row < 15:
                roles.append(Role(name=name, area=area, required=True, cost_hourly=0, source_row=row))
                continue
            source_cost = self._number(ws.cell(row, 6).value)
            required = self._text(ws.cell(row, 4).value).upper() != "NO"
            collaborators = self._number(ws.cell(row, 5).value) or 1
            roles.append(Role(
                name=name,
                area=area,
                required=required,
                collaborators=collaborators,
                source_cost=source_cost,
                cost_hourly=round(source_cost / 160, 2) if source_cost else 0,
                source_row=row,
            ))
        return roles

    def _build_index(self) -> None:
        if not self.chunks:
            return
        self.vectorizer = TfidfVectorizer(
            lowercase=True, ngram_range=(1, 2), strip_accents="unicode", max_features=12000
        )
        self.matrix = self.vectorizer.fit_transform(chunk["text"] for chunk in self.chunks)

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value)).strip()

    @staticmethod
    def _number(value: Any) -> float:
        try:
            return float(value or 0)
        except (TypeError, ValueError):
            return 0
